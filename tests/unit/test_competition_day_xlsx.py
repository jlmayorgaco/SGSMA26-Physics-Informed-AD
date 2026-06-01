from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.submission.test_day_xlsx import (
    BASE_REQUIRED_SIGNALS,
    build_metrics_rows,
    canonicalize_frame_for_inference,
    label_workbook,
    prepare_workbook_for_inference,
    process_submission,
    validate_preparation,
)


def _bus_frame(bus: int, include_label: bool) -> pd.DataFrame:
    data = {"TIMESTAMP": [0.0, 0.033]}
    for signal in BASE_REQUIRED_SIGNALS:
        column = f"BUS{bus}_{signal}"
        if signal == "Freq":
            column = f"BUS{bus}_FREQ"
        data[column] = [1.0, 2.0]
    data["DATA_PRESENT"] = [1, 1]
    if include_label:
        data["label"] = ["?", "?"]
    return pd.DataFrame(data)


def _make_test_workbook(path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _bus_frame(2, include_label=True).to_excel(writer, sheet_name="Bus2", index=False)
        _bus_frame(5, include_label=False).to_excel(writer, sheet_name="PMU Bus 5", index=False)
        pd.DataFrame({"Metric": ["Accuracy"], "Value": ["?"]}).to_excel(
            writer,
            sheet_name="Evaluation Metrics",
            index=False,
        )


def test_canonicalize_frame_for_inference_matches_training_names() -> None:
    frame = pd.DataFrame(
        {
            "TIMESTAMP": [0.0],
            "BUS2_FREQ": [59.97],
            "BUS2_ROCOF": [0.0],
            "label": ["?"],
        }
    )

    out = canonicalize_frame_for_inference(frame, bus=2)

    assert "BUS2_Freq" in out.columns
    assert "BUS2_FREQ" not in out.columns
    assert "label" not in out.columns
    assert "DATA_PRESENT" in out.columns


def test_prepare_workbook_for_inference_writes_bus_csvs(tmp_path: Path) -> None:
    workbook = tmp_path / "test.xlsx"
    _make_test_workbook(workbook)

    prep = prepare_workbook_for_inference(workbook, tmp_path / "work")

    assert prep.total_rows == 4
    assert sorted(item.bus for item in prep.bus_sheets) == [2, 5]
    bus2_csv = pd.read_csv(prep.csv_dir / "Bus2_Competition_Data_nanmask.csv")
    assert "BUS2_Freq" in bus2_csv.columns
    assert "label" not in bus2_csv.columns
    assert "DATA_PRESENT" in bus2_csv.columns


def test_label_workbook_fills_labels_and_metrics_sheet(tmp_path: Path) -> None:
    workbook = tmp_path / "test.xlsx"
    output = tmp_path / "out.xlsx"
    _make_test_workbook(workbook)
    prep = prepare_workbook_for_inference(workbook, tmp_path / "work")
    predictions = pd.DataFrame(
        {
            "TIMESTAMP": [0.0, 0.0, 0.033, 0.033],
            "Bus": [2, 5, 2, 5],
            "Predicted_Event": [0, 1, 1, 0],
            "Predicted_Location": ["none", "BUS5", "BUS2", "none"],
        }
    )
    metrics_rows = build_metrics_rows(
        predictions=predictions,
        diagnostics={"model": "test_model", "window_seconds": 30.0},
        total_rows=prep.total_rows,
        source_truth=prep.source_truth,
        model_dir=tmp_path / "models",
    )

    label_workbook(
        input_xlsx=workbook,
        output_xlsx=output,
        bus_sheets=prep.bus_sheets,
        predictions=predictions,
        metrics_rows=metrics_rows,
    )

    wb = load_workbook(output)
    assert "Evaluation Metrics" in wb.sheetnames
    bus2_label_col = wb["Bus2"].max_column
    bus5_label_col = wb["PMU Bus 5"].max_column
    assert wb["Bus2"].cell(row=1, column=bus2_label_col).value == "label"
    assert [wb["Bus2"].cell(row=i, column=bus2_label_col).value for i in range(2, 4)] == [0, 1]
    assert wb["PMU Bus 5"].cell(row=1, column=bus5_label_col).value == "label"
    assert [wb["PMU Bus 5"].cell(row=i, column=bus5_label_col).value for i in range(2, 4)] == [1, 0]
    assert wb["Evaluation Metrics"]["A1"].value == "Metric"
    assert wb["Evaluation Metrics"]["A11"].value == "Predicted abnormal rows"


def test_process_submission_keeps_upload_folder_xlsx_only(tmp_path: Path, monkeypatch) -> None:
    test1 = tmp_path / "test1.xlsx"
    test2 = tmp_path / "test2.xlsx"
    _make_test_workbook(test1)
    _make_test_workbook(test2)

    def fake_process_test_workbook(**kwargs):
        output_xlsx = kwargs["output_xlsx"]
        output_xlsx.write_bytes(b"fake")
        diagnostics_dir = kwargs["diagnostics_dir"]
        diagnostics_dir.mkdir(parents=True, exist_ok=True)
        (diagnostics_dir / f"{output_xlsx.stem}.diagnostics.json").write_text("{}")
        return {"output": str(output_xlsx)}

    monkeypatch.setattr("src.submission.test_day_xlsx.process_test_workbook", fake_process_test_workbook)
    stale_dir = tmp_path / "prepared" / "J_Mayorga_SGSMA2026"
    stale_dir.mkdir(parents=True)
    (stale_dir / "old_diagnostics.json").write_text("{}")

    process_submission(
        test_inputs=[test1, test2],
        output_dir=tmp_path / "prepared",
        expected_buses=None,
    )

    upload_dir = tmp_path / "prepared" / "J_Mayorga_SGSMA2026"
    assert sorted(path.name for path in upload_dir.iterdir()) == [
        "J_Mayorga_Results_Test1.xlsx",
        "J_Mayorga_Results_Test2.xlsx",
    ]
    assert (tmp_path / "prepared" / "_diagnostics" / "submission_summary.json").exists()


def test_validate_preparation_rejects_missing_expected_bus(tmp_path: Path) -> None:
    workbook = tmp_path / "test.xlsx"
    _make_test_workbook(workbook)
    prep = prepare_workbook_for_inference(workbook, tmp_path / "work")

    with pytest.raises(ValueError, match="missing expected bus sheets"):
        validate_preparation(prep, expected_buses=(2, 5, 6))


def test_prepare_workbook_rejects_label_not_last(tmp_path: Path) -> None:
    workbook = tmp_path / "test.xlsx"
    bad = _bus_frame(2, include_label=False)
    bad.insert(1, "label", ["?", "?"])
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        bad.to_excel(writer, sheet_name="Bus2", index=False)

    with pytest.raises(ValueError, match="label column exists but is not the last column"):
        prepare_workbook_for_inference(workbook, tmp_path / "work")
