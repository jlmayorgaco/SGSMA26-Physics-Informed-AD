"""Competition-day XLSX packaging utilities.

The organizer's day-of-test format is workbook oriented: one workbook per test
case, one sheet per bus, and a blank ``label`` column to fill. The production
runtime in ``main.py`` consumes Bus*.csv folders, so this module bridges both
formats while preserving the original workbook layout.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell import WriteOnlyCell
from sklearn.metrics import accuracy_score, f1_score

from src.models.hybrid_submission import run_hybrid_submission_prediction


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_PARTICIPANT = "J_Mayorga"
DEFAULT_MODEL_DIR = ROOT / "models"
DEFAULT_OUTPUT_DIR = ROOT / "output" / "competition_day"
DEFAULT_EXPECTED_BUSES = (2, 5, 6, 10, 19, 22, 29, 39)
BUS_RE = re.compile(r"BUS[\s_-]*([0-9]+)", re.IGNORECASE)
BUS_COLUMN_RE = re.compile(r"^BUS([0-9]+)_", re.IGNORECASE)
LABEL_COLUMN_NAMES = {"label", "predicted_event", "predicted label", "predicted_label"}
METRICS_SHEET_NAMES = {"metrics", "evaluation metrics", "evaluation_metrics"}
HEADER_FILL = PatternFill("solid", fgColor="FFF2CC")
METRICS_HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
BASE_REQUIRED_SIGNALS = (
    "VA_ANG",
    "VA_MAG",
    "VB_ANG",
    "VB_MAG",
    "VC_ANG",
    "VC_MAG",
    "IA_ANG",
    "IA_MAG",
    "IB_ANG",
    "IB_MAG",
    "IC_ANG",
    "IC_MAG",
    "Freq",
    "ROCOF",
)


@dataclass(frozen=True)
class PreparedBusSheet:
    sheet_name: str
    bus: int
    rows: int
    had_label_column: bool


@dataclass(frozen=True)
class WorkbookPreparation:
    csv_dir: Path
    bus_sheets: list[PreparedBusSheet]
    total_rows: int
    source_truth: dict[tuple[int, tuple[str, Any]], int] | None
    warnings: list[str]


@dataclass(frozen=True)
class CsvBusSource:
    sheet_name: str
    bus: int
    source_name: str
    original_columns: list[str]
    model_csv_path: Path
    rows: int


@dataclass(frozen=True)
class CsvPreparation:
    csv_dir: Path
    bus_sources: list[CsvBusSource]
    total_rows: int
    source_truth: dict[tuple[int, tuple[str, Any]], int] | None
    warnings: list[str]


def _timestamp_key(value: Any) -> tuple[str, Any]:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return ("text", "" if value is None else str(value).strip())
    if not np.isfinite(numeric):
        return ("text", "" if value is None else str(value).strip())
    return ("number", round(numeric, 9))


def _normalize_label_column_name(column: object) -> str:
    return str(column).strip().lower()


def _is_label_column(column: object) -> bool:
    return _normalize_label_column_name(column) in LABEL_COLUMN_NAMES


def _is_metrics_sheet(sheet_name: str) -> bool:
    return sheet_name.strip().lower() in METRICS_SHEET_NAMES


def _column_key(column: object) -> str:
    return str(column).strip().upper()


def _metadata_column_name(column: object) -> str | None:
    key = _column_key(column)
    if key == "TIMESTAMP":
        return "TIMESTAMP"
    if key == "DATA_PRESENT":
        return "DATA_PRESENT"
    if key == "EVENT":
        return "Event"
    return None


def _infer_bus_from_columns(columns: list[object]) -> int | None:
    buses = set()
    for column in columns:
        match = BUS_COLUMN_RE.match(str(column).strip())
        if match:
            buses.add(int(match.group(1)))
    return buses.pop() if len(buses) == 1 else None


def infer_bus_from_sheet(sheet_name: str, columns: list[object]) -> int | None:
    match = BUS_RE.search(sheet_name)
    if match:
        return int(match.group(1))
    return _infer_bus_from_columns(columns)


def sheet_name_for_bus(bus: int) -> str:
    return f"Bus{int(bus)}"


def _canonical_signal_suffix(suffix: str) -> str:
    suffix_key = suffix.strip().upper()
    if suffix_key == "FREQ":
        return "Freq"
    return suffix_key


def _canonical_inference_column(column: object, bus: int) -> str:
    name = str(column).strip()
    metadata_name = _metadata_column_name(name)
    if metadata_name is not None:
        return metadata_name
    match = BUS_COLUMN_RE.match(name)
    if match and int(match.group(1)) == int(bus):
        suffix = name[match.end() :]
        return f"BUS{bus}_{_canonical_signal_suffix(suffix)}"
    return name


def canonicalize_frame_for_inference(frame: pd.DataFrame, bus: int) -> pd.DataFrame:
    """Return a model-facing copy without blank label columns.

    The workbook itself is not modified here. This only creates the temporary
    Bus*.csv files consumed by the existing inference runtime.
    """

    keep_columns = [column for column in frame.columns if not _is_label_column(column)]
    out = frame.loc[:, keep_columns].copy()
    out.columns = [_canonical_inference_column(column, bus) for column in out.columns]
    if "TIMESTAMP" not in out.columns:
        raise ValueError(f"Bus {bus} sheet is missing TIMESTAMP")
    if "DATA_PRESENT" not in out.columns:
        out["DATA_PRESENT"] = 1
    return out


def _expected_columns_for_bus(bus: int) -> set[str]:
    return {f"BUS{bus}_{_canonical_signal_suffix(signal)}".upper() for signal in BASE_REQUIRED_SIGNALS}


def _validate_bus_frame(frame: pd.DataFrame, sheet_name: str, bus: int) -> list[str]:
    warnings: list[str] = []
    canonical_columns = [_canonical_inference_column(column, bus) for column in frame.columns]
    column_keys = {_column_key(column) for column in canonical_columns}
    if "TIMESTAMP" not in column_keys:
        raise ValueError(f"{sheet_name}: missing required TIMESTAMP column")
    missing = sorted(_expected_columns_for_bus(bus) - column_keys)
    if missing:
        raise ValueError(f"{sheet_name}: missing required bus columns: {', '.join(missing)}")
    if "DATA_PRESENT" not in column_keys:
        warnings.append(f"{sheet_name}: DATA_PRESENT missing; temporary inference CSV will assume DATA_PRESENT=1")

    label_positions = [idx for idx, column in enumerate(frame.columns) if _is_label_column(column)]
    if label_positions and label_positions[-1] != len(frame.columns) - 1:
        raise ValueError(f"{sheet_name}: label column exists but is not the last column")
    if not label_positions:
        warnings.append(f"{sheet_name}: label column missing; it will be appended as the last column")

    timestamp_col = next(column for column in frame.columns if _metadata_column_name(column) == "TIMESTAMP")
    keys = [_timestamp_key(value) for value in frame[timestamp_col].tolist()]
    if len(keys) != len(set(keys)):
        raise ValueError(f"{sheet_name}: duplicate TIMESTAMP values detected")
    if frame[timestamp_col].isna().any():
        raise ValueError(f"{sheet_name}: blank TIMESTAMP values detected")
    return warnings


def _label_truth_from_frame(frame: pd.DataFrame) -> pd.Series | None:
    candidates = [column for column in frame.columns if str(column).strip() == "Event"]
    for column in candidates:
        values = pd.to_numeric(frame[column], errors="coerce")
        if values.notna().any():
            return values
    return None


def prepare_workbook_for_inference(input_xlsx: Path, work_dir: Path) -> WorkbookPreparation:
    input_xlsx = Path(input_xlsx)
    if input_xlsx.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"Expected .xlsx or .xlsm input, got {input_xlsx}")
    if not input_xlsx.exists():
        raise FileNotFoundError(input_xlsx)

    csv_dir = work_dir / "csv_input"
    if csv_dir.exists():
        shutil.rmtree(csv_dir)
    csv_dir.mkdir(parents=True, exist_ok=True)

    xls = pd.ExcelFile(input_xlsx)
    prepared: list[PreparedBusSheet] = []
    truth_lookup: dict[tuple[int, tuple[str, Any]], int] = {}
    warnings: list[str] = []
    seen_buses: dict[int, str] = {}
    for sheet_name in xls.sheet_names:
        if _is_metrics_sheet(sheet_name):
            continue
        frame = pd.read_excel(xls, sheet_name=sheet_name)
        if frame.empty or "TIMESTAMP" not in [_column_key(column) for column in frame.columns]:
            continue
        frame.columns = [str(column).strip() for column in frame.columns]
        bus = infer_bus_from_sheet(sheet_name, list(frame.columns))
        if bus is None:
            continue
        if int(bus) in seen_buses:
            raise ValueError(f"Duplicate bus {bus} sheets detected: {seen_buses[int(bus)]} and {sheet_name}")
        seen_buses[int(bus)] = sheet_name
        warnings.extend(_validate_bus_frame(frame, sheet_name, int(bus)))
        model_frame = canonicalize_frame_for_inference(frame, bus)
        model_frame.to_csv(csv_dir / f"Bus{bus}_Competition_Data_nanmask.csv", index=False)
        had_label_column = any(_is_label_column(column) for column in frame.columns)
        truth = _label_truth_from_frame(frame)
        if truth is not None:
            for timestamp, value in zip(frame["TIMESTAMP"], truth, strict=False):
                if pd.notna(value):
                    truth_lookup[(int(bus), _timestamp_key(timestamp))] = int(value)
        prepared.append(
            PreparedBusSheet(
                sheet_name=sheet_name,
                bus=int(bus),
                rows=int(len(frame)),
                had_label_column=had_label_column,
            )
        )

    if not prepared:
        raise ValueError(f"No bus sheets were detected in {input_xlsx}")

    total_rows = int(sum(item.rows for item in prepared))
    source_truth = truth_lookup if truth_lookup else None
    return WorkbookPreparation(
        csv_dir=csv_dir,
        bus_sheets=prepared,
        total_rows=total_rows,
        source_truth=source_truth,
        warnings=warnings,
    )


def _safe_extract_zip_csvs(zip_path: Path, extract_dir: Path) -> list[Path]:
    if extract_dir.exists():
        shutil.rmtree(extract_dir)
    extract_dir.mkdir(parents=True, exist_ok=True)
    csv_paths: list[Path] = []
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            if info.is_dir() or not info.filename.lower().endswith(".csv"):
                continue
            name = Path(info.filename).name
            if not name:
                continue
            target = extract_dir / name
            with zf.open(info) as src, target.open("wb") as dst:
                shutil.copyfileobj(src, dst)
            csv_paths.append(target)
    return sorted(csv_paths)


def _csv_paths_from_input(input_path: Path, work_dir: Path) -> list[Path]:
    input_path = Path(input_path)
    if input_path.suffix.lower() == ".zip":
        return _safe_extract_zip_csvs(input_path, work_dir / "zip_csvs")
    if input_path.is_dir():
        return sorted(input_path.rglob("*.csv"))
    if input_path.suffix.lower() == ".csv":
        return [input_path]
    raise ValueError(f"Expected .xlsx, .xlsm, .zip, .csv, or CSV folder input, got {input_path}")


def prepare_csv_input_for_inference(input_path: Path, work_dir: Path) -> CsvPreparation:
    input_path = Path(input_path)
    if not input_path.exists():
        raise FileNotFoundError(input_path)
    csv_dir = work_dir / "csv_input"
    if csv_dir.exists():
        shutil.rmtree(csv_dir)
    csv_dir.mkdir(parents=True, exist_ok=True)

    warnings: list[str] = []
    truth_lookup: dict[tuple[int, tuple[str, Any]], int] = {}
    bus_sources: list[CsvBusSource] = []
    seen_buses: dict[int, str] = {}

    for csv_path in _csv_paths_from_input(input_path, work_dir):
        frame = pd.read_csv(csv_path)
        frame.columns = [str(column).strip() for column in frame.columns]
        if frame.empty or "TIMESTAMP" not in [_column_key(column) for column in frame.columns]:
            continue
        bus = infer_bus_from_sheet(csv_path.stem, list(frame.columns))
        if bus is None:
            continue
        if int(bus) in seen_buses:
            raise ValueError(f"Duplicate bus {bus} CSV files detected: {seen_buses[int(bus)]} and {csv_path.name}")
        seen_buses[int(bus)] = csv_path.name
        warnings.extend(_validate_bus_frame(frame, csv_path.name, int(bus)))
        model_frame = canonicalize_frame_for_inference(frame, int(bus))
        model_csv_path = csv_dir / f"Bus{bus}_Competition_Data_nanmask.csv"
        model_frame.to_csv(model_csv_path, index=False)
        truth = _label_truth_from_frame(frame)
        if truth is not None:
            for timestamp, value in zip(frame["TIMESTAMP"], truth, strict=False):
                if pd.notna(value):
                    truth_lookup[(int(bus), _timestamp_key(timestamp))] = int(value)
        bus_sources.append(
            CsvBusSource(
                sheet_name=sheet_name_for_bus(int(bus)),
                bus=int(bus),
                source_name=csv_path.name,
                original_columns=list(frame.columns),
                model_csv_path=model_csv_path,
                rows=int(len(frame)),
            )
        )

    if not bus_sources:
        raise ValueError(f"No bus CSV files were detected in {input_path}")
    return CsvPreparation(
        csv_dir=csv_dir,
        bus_sources=sorted(bus_sources, key=lambda item: item.bus),
        total_rows=int(sum(item.rows for item in bus_sources)),
        source_truth=truth_lookup if truth_lookup else None,
        warnings=warnings,
    )


def validate_preparation(prep: WorkbookPreparation, expected_buses: tuple[int, ...] | None) -> None:
    if expected_buses is None:
        return
    actual = {item.bus for item in prep.bus_sheets}
    expected = set(expected_buses)
    missing = sorted(expected - actual)
    extra = sorted(actual - expected)
    errors = []
    if missing:
        errors.append(f"missing expected bus sheets: {missing}")
    if extra:
        errors.append(f"unexpected bus sheets: {extra}")
    if errors:
        raise ValueError("; ".join(errors))


def validate_csv_preparation(prep: CsvPreparation, expected_buses: tuple[int, ...] | None) -> None:
    if expected_buses is None:
        return
    actual = {item.bus for item in prep.bus_sources}
    expected = set(expected_buses)
    missing = sorted(expected - actual)
    extra = sorted(actual - expected)
    errors = []
    if missing:
        errors.append(f"missing expected bus CSVs: {missing}")
    if extra:
        errors.append(f"unexpected bus CSVs: {extra}")
    if errors:
        raise ValueError("; ".join(errors))


def _prediction_lookup(predictions: pd.DataFrame) -> dict[tuple[int, tuple[str, Any]], int]:
    lookup: dict[tuple[int, tuple[str, Any]], int] = {}
    for row in predictions.itertuples(index=False):
        lookup[(int(row.Bus), _timestamp_key(row.TIMESTAMP))] = int(row.Predicted_Event)
    return lookup


def _header_values(ws) -> list[str]:
    return ["" if cell.value is None else str(cell.value).strip() for cell in ws[1]]


def _label_column_index(ws) -> int:
    headers = _header_values(ws)
    for idx, header in enumerate(headers, start=1):
        if _is_label_column(header):
            ws.cell(row=1, column=idx).value = "label"
            return idx
    idx = ws.max_column + 1
    ws.cell(row=1, column=idx).value = "label"
    return idx


def _timestamp_column_index(ws) -> int:
    for idx, header in enumerate(_header_values(ws), start=1):
        if _metadata_column_name(header) == "TIMESTAMP":
            return idx
    raise ValueError(f"Sheet {ws.title} is missing TIMESTAMP")


def _write_metrics_sheet(wb, rows: list[tuple[str, Any]]) -> None:
    existing = None
    for name in wb.sheetnames:
        if _is_metrics_sheet(name):
            existing = name
            break
    if existing is None:
        ws = wb.create_sheet("Evaluation Metrics")
    else:
        ws = wb[existing]
        ws.delete_rows(1, ws.max_row)
        ws.title = "Evaluation Metrics"

    ws["A1"] = "Metric"
    ws["B1"] = "Value"
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = METRICS_HEADER_FILL

    for row_idx, (metric, value) in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1).value = metric
        ws.cell(row=row_idx, column=2).value = value
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 78


def _write_metrics_sheet_write_only(wb: Workbook, rows: list[tuple[str, Any]]) -> None:
    ws = wb.create_sheet("Evaluation Metrics")
    header = []
    for value in ("Metric", "Value"):
        cell = WriteOnlyCell(ws, value=value)
        cell.font = Font(bold=True)
        cell.fill = METRICS_HEADER_FILL
        header.append(cell)
    ws.append(header)
    for metric, value in rows:
        ws.append([metric, value])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 78


def _excel_cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if stripped == "":
            return None
        try:
            numeric = float(stripped)
        except ValueError:
            return value
        return numeric if np.isfinite(numeric) else value
    return value


def label_workbook(
    input_xlsx: Path,
    output_xlsx: Path,
    bus_sheets: list[PreparedBusSheet],
    predictions: pd.DataFrame,
    metrics_rows: list[tuple[str, Any]],
) -> None:
    wb = load_workbook(input_xlsx)
    lookup = _prediction_lookup(predictions)
    missing: list[str] = []

    for item in bus_sheets:
        ws = wb[item.sheet_name]
        timestamp_col = _timestamp_column_index(ws)
        label_col = _label_column_index(ws)
        ws.cell(row=1, column=label_col).fill = HEADER_FILL
        ws.cell(row=1, column=label_col).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(label_col)].width = 12
        for row_idx in range(2, ws.max_row + 1):
            timestamp = ws.cell(row=row_idx, column=timestamp_col).value
            key = (int(item.bus), _timestamp_key(timestamp))
            if key not in lookup:
                missing.append(f"{item.sheet_name}!row{row_idx}")
                continue
            label = int(lookup[key])
            ws.cell(row=row_idx, column=label_col).value = label

    if missing:
        sample = ", ".join(missing[:10])
        raise ValueError(f"Missing predictions for {len(missing)} workbook rows; first rows: {sample}")

    _write_metrics_sheet(wb, metrics_rows)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)


def create_labeled_workbook_from_csvs(
    output_xlsx: Path,
    bus_sources: list[CsvBusSource],
    predictions: pd.DataFrame,
    metrics_rows: list[tuple[str, Any]],
) -> None:
    lookup = _prediction_lookup(predictions)
    wb = Workbook(write_only=True)
    if wb.worksheets:
        wb.remove(wb.worksheets[0])
    missing: list[str] = []

    for source in sorted(bus_sources, key=lambda item: item.bus):
        ws = wb.create_sheet(source.sheet_name)
        header_values = list(source.original_columns)
        if not header_values or not _is_label_column(header_values[-1]):
            header_values.append("label")
        else:
            header_values[-1] = "label"
        header = []
        for value in header_values:
            cell = WriteOnlyCell(ws, value=value)
            if _is_label_column(value):
                cell.font = Font(bold=True)
                cell.fill = HEADER_FILL
            header.append(cell)
        ws.append(header)

        with source.model_csv_path.open("r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            for row_index, row in enumerate(reader, start=2):
                timestamp = row.get("TIMESTAMP")
                key = (int(source.bus), _timestamp_key(timestamp))
                if key not in lookup:
                    missing.append(f"{source.sheet_name}!row{row_index}")
                    label = None
                else:
                    label = int(lookup[key])
                output_row = [
                    _excel_cell_value(row.get(_canonical_inference_column(column, source.bus), row.get(column)))
                    for column in source.original_columns
                ]
                has_label_last = bool(source.original_columns and _is_label_column(source.original_columns[-1]))
                if has_label_last:
                    output_row[-1] = label
                else:
                    output_row.append(label)
                ws.append(output_row)

    if missing:
        sample = ", ".join(missing[:10])
        raise ValueError(f"Missing predictions for {len(missing)} CSV rows; first rows: {sample}")

    _write_metrics_sheet_write_only(wb, metrics_rows)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)


def _model_size_mb(model_dir: Path) -> float:
    if not model_dir.exists():
        return 0.0
    total = sum(path.stat().st_size for path in model_dir.rglob("*") if path.is_file())
    return round(total / (1024 * 1024), 2)


def _load_feature_count(model_dir: Path) -> int | None:
    metrics_path = model_dir / "final_metrics.json"
    if metrics_path.exists():
        try:
            payload = json.loads(metrics_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            payload = {}
        selected = payload.get("selected", {})
        if "n_features" in selected:
            return int(selected["n_features"])
    feature_path = model_dir / "feature_columns.json"
    if feature_path.exists():
        try:
            return len(json.loads(feature_path.read_text(encoding="utf-8")))
        except json.JSONDecodeError:
            return None
    return None


def _distribution(predictions: pd.DataFrame) -> dict[str, int]:
    counts = predictions["Predicted_Event"].astype(int).value_counts().sort_index()
    return {str(int(label)): int(count) for label, count in counts.items()}


def _jsonable_bus_source(source: CsvBusSource) -> dict[str, Any]:
    return {
        "sheet_name": source.sheet_name,
        "bus": int(source.bus),
        "source_name": source.source_name,
        "original_columns": list(source.original_columns),
        "model_csv_path": str(source.model_csv_path),
        "rows": int(source.rows),
    }


def _maybe_metrics_from_truth(
    truth: dict[tuple[int, tuple[str, Any]], int] | None,
    predictions: pd.DataFrame,
) -> dict[str, float] | None:
    if truth is None:
        return None
    y_true: list[int] = []
    y_pred: list[int] = []
    for row in predictions.itertuples(index=False):
        key = (int(row.Bus), _timestamp_key(row.TIMESTAMP))
        if key not in truth:
            return None
        y_true.append(int(truth[key]))
        y_pred.append(int(row.Predicted_Event))
    if len(y_true) != len(predictions) or not y_true:
        return None
    y_true_arr = np.asarray(y_true, dtype=int)
    y_pred_arr = np.asarray(y_pred, dtype=int)
    return {
        "Accuracy": float(accuracy_score(y_true_arr, y_pred_arr)),
        "F1 Score (weighted)": float(f1_score(y_true_arr, y_pred_arr, average="weighted", zero_division=0)),
        "F1 Score (macro)": float(f1_score(y_true_arr, y_pred_arr, average="macro", zero_division=0)),
        "Detection F1 (binary normal vs abnormal)": float(
            f1_score((y_true_arr != 0).astype(int), (y_pred_arr != 0).astype(int), zero_division=0)
        ),
    }


def build_metrics_rows(
    predictions: pd.DataFrame,
    diagnostics: dict[str, Any],
    total_rows: int,
    source_truth: dict[tuple[int, tuple[str, Any]], int] | None,
    model_dir: Path,
) -> list[tuple[str, Any]]:
    computed = _maybe_metrics_from_truth(source_truth, predictions)
    unavailable = "N/A - hidden ground truth labels unavailable"
    distribution = _distribution(predictions)
    feature_count = _load_feature_count(model_dir)
    trainable = "N/A (ExtraTrees ensemble; non-neural model)"
    if feature_count is not None:
        trainable = f"N/A (ExtraTrees ensemble; {feature_count:,} feature inputs)"

    return [
        ("Accuracy", computed["Accuracy"] if computed else unavailable),
        ("F1 Score (weighted)", computed["F1 Score (weighted)"] if computed else unavailable),
        ("F1 Score (macro)", computed["F1 Score (macro)"] if computed else unavailable),
        (
            "Detection F1 (binary normal vs abnormal)",
            computed["Detection F1 (binary normal vs abnormal)"] if computed else unavailable,
        ),
        ("Model", diagnostics.get("model", "sgms_extra_trees_windowed_v2")),
        ("Trainable parameters", trainable),
        ("Model size", f"{_model_size_mb(model_dir):.2f} MB"),
        ("Window length", f"{diagnostics.get('window_seconds', 30.0)} s"),
        ("Test rows", int(total_rows)),
        ("Predicted abnormal rows", int((predictions["Predicted_Event"].astype(int) != 0).sum())),
        ("Predicted label distribution", json.dumps(distribution, sort_keys=True)),
    ]


def process_test_workbook(
    input_xlsx: Path,
    output_xlsx: Path,
    work_dir: Path,
    diagnostics_dir: Path,
    model_dir: Path = DEFAULT_MODEL_DIR,
    window_seconds: float = 30.0,
    expected_buses: tuple[int, ...] | None = DEFAULT_EXPECTED_BUSES,
) -> dict[str, Any]:
    prep = prepare_workbook_for_inference(input_xlsx=input_xlsx, work_dir=work_dir)
    validate_preparation(prep, expected_buses=expected_buses)
    predictions, diagnostics = run_hybrid_submission_prediction(
        input_dir=prep.csv_dir,
        output_csv=work_dir / "predictions.csv",
        ml_model_dir=model_dir,
        ml_window_seconds=window_seconds,
    )
    metrics_rows = build_metrics_rows(
        predictions=predictions,
        diagnostics=diagnostics,
        total_rows=prep.total_rows,
        source_truth=prep.source_truth,
        model_dir=model_dir,
    )
    label_workbook(
        input_xlsx=input_xlsx,
        output_xlsx=output_xlsx,
        bus_sheets=prep.bus_sheets,
        predictions=predictions,
        metrics_rows=metrics_rows,
    )
    summary = {
        "input": str(input_xlsx.resolve()),
        "output": str(output_xlsx.resolve()),
        "bus_sheets": [item.__dict__ for item in prep.bus_sheets],
        "total_rows": prep.total_rows,
        "validation_warnings": prep.warnings,
        "predicted_label_distribution": _distribution(predictions),
        "diagnostics": diagnostics,
    }
    diagnostics_dir.mkdir(parents=True, exist_ok=True)
    (diagnostics_dir / f"{output_xlsx.stem}.diagnostics.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    return summary


def process_test_csv_input(
    input_path: Path,
    output_xlsx: Path,
    work_dir: Path,
    diagnostics_dir: Path,
    model_dir: Path = DEFAULT_MODEL_DIR,
    window_seconds: float = 30.0,
    expected_buses: tuple[int, ...] | None = DEFAULT_EXPECTED_BUSES,
) -> dict[str, Any]:
    prep = prepare_csv_input_for_inference(input_path=input_path, work_dir=work_dir)
    validate_csv_preparation(prep, expected_buses=expected_buses)
    predictions, diagnostics = run_hybrid_submission_prediction(
        input_dir=prep.csv_dir,
        output_csv=work_dir / "predictions.csv",
        ml_model_dir=model_dir,
        ml_window_seconds=window_seconds,
    )
    metrics_rows = build_metrics_rows(
        predictions=predictions,
        diagnostics=diagnostics,
        total_rows=prep.total_rows,
        source_truth=prep.source_truth,
        model_dir=model_dir,
    )
    create_labeled_workbook_from_csvs(
        output_xlsx=output_xlsx,
        bus_sources=prep.bus_sources,
        predictions=predictions,
        metrics_rows=metrics_rows,
    )
    summary = {
        "input": str(Path(input_path).resolve()),
        "output": str(output_xlsx.resolve()),
        "bus_sheets": [_jsonable_bus_source(item) for item in prep.bus_sources],
        "total_rows": prep.total_rows,
        "validation_warnings": prep.warnings,
        "predicted_label_distribution": _distribution(predictions),
        "diagnostics": diagnostics,
    }
    diagnostics_dir.mkdir(parents=True, exist_ok=True)
    (diagnostics_dir / f"{output_xlsx.stem}.diagnostics.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    return summary


def process_test_input(
    input_path: Path,
    output_xlsx: Path,
    work_dir: Path,
    diagnostics_dir: Path,
    model_dir: Path = DEFAULT_MODEL_DIR,
    window_seconds: float = 30.0,
    expected_buses: tuple[int, ...] | None = DEFAULT_EXPECTED_BUSES,
) -> dict[str, Any]:
    suffix = Path(input_path).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return process_test_workbook(
            input_xlsx=input_path,
            output_xlsx=output_xlsx,
            work_dir=work_dir,
            diagnostics_dir=diagnostics_dir,
            model_dir=model_dir,
            window_seconds=window_seconds,
            expected_buses=expected_buses,
        )
    return process_test_csv_input(
        input_path=input_path,
        output_xlsx=output_xlsx,
        work_dir=work_dir,
        diagnostics_dir=diagnostics_dir,
        model_dir=model_dir,
        window_seconds=window_seconds,
        expected_buses=expected_buses,
    )


def process_submission(
    test_inputs: list[Path],
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    participant: str = DEFAULT_PARTICIPANT,
    model_dir: Path = DEFAULT_MODEL_DIR,
    window_seconds: float = 30.0,
    keep_work: bool = False,
    expected_buses: tuple[int, ...] | None = DEFAULT_EXPECTED_BUSES,
) -> list[dict[str, Any]]:
    if len(test_inputs) != 2:
        raise ValueError("Exactly two test workbooks are required.")
    submission_dir = output_dir / f"{participant}_SGSMA2026"
    work_root = output_dir / "_work"
    diagnostics_dir = output_dir / "_diagnostics"
    if submission_dir.exists():
        shutil.rmtree(submission_dir)
    if work_root.exists() and not keep_work:
        shutil.rmtree(work_root)
    if diagnostics_dir.exists() and not keep_work:
        shutil.rmtree(diagnostics_dir)
    submission_dir.mkdir(parents=True, exist_ok=True)
    work_root.mkdir(parents=True, exist_ok=True)

    summaries = []
    for index, input_xlsx in enumerate(test_inputs, start=1):
        output_xlsx = submission_dir / f"{participant}_Results_Test{index}.xlsx"
        summary = process_test_input(
            input_path=Path(input_xlsx),
            output_xlsx=output_xlsx,
            work_dir=work_root / f"test{index}",
            diagnostics_dir=diagnostics_dir,
            model_dir=model_dir,
            window_seconds=window_seconds,
            expected_buses=expected_buses,
        )
        summaries.append(summary)
    diagnostics_dir.mkdir(parents=True, exist_ok=True)
    (diagnostics_dir / "submission_summary.json").write_text(json.dumps(summaries, indent=2), encoding="utf-8")
    return summaries


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Prepare SGSMA 2026 competition-day XLSX submission files.")
    parser.add_argument("inputs", nargs="*", type=Path, help="Two input XLSX test workbooks, in Test1/Test2 order.")
    parser.add_argument("--test1", type=Path, default=None, help="Input workbook for test case 1.")
    parser.add_argument("--test2", type=Path, default=None, help="Input workbook for test case 2.")
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUTPUT_DIR, help="Output parent directory.")
    parser.add_argument("--participant", default=DEFAULT_PARTICIPANT, help="Participant prefix, e.g. J_Mayorga.")
    parser.add_argument("--model-dir", type=Path, default=DEFAULT_MODEL_DIR, help="Validated model bundle directory.")
    parser.add_argument("--window-seconds", type=float, default=30.0, help="Inference window length in seconds.")
    parser.add_argument("--keep-work", action="store_true", help="Keep intermediate CSV folders under output/_work.")
    return parser.parse_args()


def main() -> None:
    args = _parse_args()
    inputs = list(args.inputs)
    if args.test1 is not None or args.test2 is not None:
        if args.test1 is None or args.test2 is None:
            raise SystemExit("Both --test1 and --test2 must be provided together.")
        inputs = [args.test1, args.test2]
    summaries = process_submission(
        test_inputs=inputs,
        output_dir=args.out_dir,
        participant=args.participant,
        model_dir=args.model_dir,
        window_seconds=args.window_seconds,
        keep_work=args.keep_work,
    )
    print(json.dumps({"outputs": [item["output"] for item in summaries]}, indent=2))


if __name__ == "__main__":
    main()
