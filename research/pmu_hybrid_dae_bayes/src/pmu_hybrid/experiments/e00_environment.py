"""Phase 1: audit the executable environment before any estimator work."""

from __future__ import annotations

import argparse
from importlib import metadata
import json
import os
from pathlib import Path
import platform
import subprocess
import sys
from typing import Any

from pmu_hybrid.utils.manifests import base_manifest, write_json


PACKAGES = (
    "numpy", "scipy", "pandas", "matplotlib", "scikit-learn", "andes",
    "pandapower", "pyarrow", "pytest", "pydantic",
)


def package_versions() -> dict[str, str]:
    """Return installed package versions, preserving an honest absence marker."""
    versions: dict[str, str] = {"python": platform.python_version()}
    for package in PACKAGES:
        try:
            versions[package] = metadata.version(package)
        except metadata.PackageNotFoundError:
            versions[package] = "NOT_INSTALLED"
    return versions


def inspect_andes_case() -> dict[str, Any]:
    """Inspect the installed native IEEE-39 workbook without executing a solve."""
    try:
        import andes
        import openpyxl
    except ImportError as exc:
        return {"status": "FAIL", "reason": f"{type(exc).__name__}: {exc}"}

    case_path = Path(andes.get_case("ieee39/ieee39_full.xlsx"))
    workbook = openpyxl.load_workbook(case_path, read_only=True, data_only=True)
    model_rows = {
        name: max(0, workbook[name].max_row - 1)
        for name in workbook.sheetnames
    }
    return {
        "status": "PASS",
        "andes_version": andes.__version__,
        "case_path": str(case_path),
        "sheets": workbook.sheetnames,
        "model_rows": model_rows,
        "dynamic_models": {
            name: model_rows.get(name, 0)
            for name in ("GENROU", "TGOV1N", "IEEEX1", "IEEEST", "BusFreq")
        },
    }


def run_andes_selftest(output_root: Path) -> dict[str, Any]:
    """Run ANDES' quick self-test with a writable, campaign-owned temp path."""
    temporary = output_root / "selftest_tmp"
    temporary.mkdir(parents=True, exist_ok=True)
    environment = os.environ.copy()
    environment["TEMP"] = str(temporary)
    environment["TMP"] = str(temporary)
    try:
        result = subprocess.run(
            [sys.executable, "-m", "andes", "selftest", "-q"],
            capture_output=True, text=True, timeout=300, env=environment,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {"status": "FAIL", "reason": f"{type(exc).__name__}: {exc}"}
    return {
        "status": "PASS" if result.returncode == 0 else "FAIL",
        "returncode": result.returncode,
        "stdout_tail": result.stdout[-4000:],
        "stderr_tail": result.stderr[-4000:],
    }


def run_andes_minimal_case() -> dict[str, Any]:
    """Run the smallest relevant shipped dynamic case without CLI log setup.

    This is the scientific environment gate.  It also gives an auditable
    fallback when a platform sandbox prevents ANDES' CLI self-test from
    creating its own log file before tests begin.
    """
    try:
        os.environ.setdefault("SYMPY_GROUND_TYPES", "python")
        import andes

        system = andes.load(
            andes.get_case("ieee39/ieee39_full.xlsx"),
            setup=False,
            no_output=True,
        )
        system.setup()
        solved = system.PFlow.run()
        if solved is False:
            return {"status": "FAIL", "reason": "ANDES PFlow returned False"}
        return {
            "status": "PASS",
            "bus_count": int(system.Bus.n),
            "line_count": int(system.Line.n),
            "generator_count": int(system.GENROU.n),
        }
    except Exception as exc:
        return {"status": "FAIL", "reason": f"{type(exc).__name__}: {exc}"}


def inspect_pandapower_case() -> dict[str, Any]:
    """Load and solve pandapower's independent AC case39 source."""
    try:
        import pandapower as pp
        import pandapower.networks as networks

        net = networks.case39()
        pp.runpp(net, calculate_voltage_angles=True)
    except Exception as exc:  # recorded as a failed gate, never hidden
        return {"status": "FAIL", "reason": f"{type(exc).__name__}: {exc}"}
    return {
        "status": "PASS" if bool(net.converged) else "FAIL",
        "pandapower_version": pp.__version__,
        "bus_count": int(len(net.bus)),
        "line_count": int(len(net.line)),
        "trafo_count": int(len(net.trafo)),
        "generator_count": int(len(net.gen) + len(net.ext_grid)),
        "system_base_mva": float(net.sn_mva),
    }


def markdown_report(audit: dict[str, Any]) -> str:
    """Render a compact, reviewable audit rather than a console-only claim."""
    versions = audit["versions"]
    andes = audit["andes_case"]
    pandapower = audit["pandapower_case"]
    lines = [
        "# Environment audit",
        "",
        f"- Python: `{versions['python']}`",
        f"- ANDES: `{versions['andes']}`",
        f"- pandapower: `{versions['pandapower']}`",
        f"- Code commit: `{audit['code_commit']}`",
        f"- ANDES self-test: **{audit['andes_selftest']['status']}**",
        f"- ANDES IEEE-39 power flow: **{audit['andes_minimal_case']['status']}**",
        f"- ANDES IEEE-39 inventory: **{andes['status']}**",
        f"- pandapower case39 AC solve: **{pandapower['status']}**",
        "",
        "## Package versions",
        "",
    ]
    lines.extend(f"- `{name}`: `{version}`" for name, version in versions.items())
    lines += ["", "## ANDES IEEE-39 case"]
    if andes["status"] == "PASS":
        lines += [
            f"- Source: `{andes['case_path']}`",
            f"- Dynamic rows: `{json.dumps(andes['dynamic_models'], sort_keys=True)}`",
        ]
    else:
        lines.append(f"- Failure: `{andes['reason']}`")
    lines += ["", "## Independent AC case"]
    if pandapower["status"] == "PASS":
        lines.append(
            "- `case39()` loaded and converged with "
            f"{pandapower['bus_count']} buses, {pandapower['line_count']} lines, "
            f"and {pandapower['trafo_count']} transformers."
        )
    else:
        lines.append(f"- Failure: `{pandapower['reason']}`")
    lines += [
        "",
        "The CLI self-test writes an ANDES log before it can execute. If the platform "
        "blocks that write, this is retained as an environment limitation, while the "
        "native IEEE-39 PFlow remains the required executable model check.",
        "",
    ]
    return "\n".join(lines)


def run(root: Path) -> dict[str, Any]:
    root = root.resolve()
    repository_root = root.parents[1]
    reports = root / "output" / "reports"
    manifests = root / "output" / "manifests"
    reports.mkdir(parents=True, exist_ok=True)
    audit = {
        **base_manifest(repository_root, seed=20260911),
        "phase": "E00",
        "versions": package_versions(),
        "andes_case": inspect_andes_case(),
        "andes_selftest": run_andes_selftest(reports),
        "andes_minimal_case": run_andes_minimal_case(),
        "pandapower_case": inspect_pandapower_case(),
    }
    audit["status"] = "PASS" if all(
        item.get("status") == "PASS"
        for item in (audit["andes_case"], audit["andes_minimal_case"], audit["pandapower_case"])
    ) else "FAIL"
    write_json(manifests / "environment_audit.json", audit)
    (reports / "environment_audit.md").write_text(markdown_report(audit), encoding="utf-8")
    return audit


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    args = parser.parse_args()
    audit = run(args.root)
    print(json.dumps({"status": audit["status"], "manifest": "output/manifests/environment_audit.json"}))
    if audit["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
